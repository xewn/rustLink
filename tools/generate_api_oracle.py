"""Exercise the original FastAPI app in an isolated temporary config/log directory."""
import ast
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import types

source=Path(sys.argv[1]) if len(sys.argv)>1 else Path('D:/PLCLink/FactoryLink')
sys.dont_write_bytecode=True
sys.path.insert(0,str(source/'backend'))
output=Path('tests/fixtures'); output.mkdir(parents=True,exist_ok=True)

def isolated_module(name, singleton_names):
    path=source/'backend'/f'{name}.py'
    tree=ast.parse(path.read_text(encoding='utf-8'))
    tree.body=[n for n in tree.body if not (isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id in singleton_names for t in n.targets))]
    module=types.ModuleType(name);module.__file__=str(path);sys.modules[name]=module
    exec(compile(tree,str(path),'exec'),module.__dict__)
    return module

with tempfile.TemporaryDirectory(prefix='factorylink-oracle-') as temp:
    root=Path(temp)
    config=isolated_module('config',{'config_manager'})
    config.get_config_dir=lambda:root
    config.config_manager=config.ConfigManager()
    logger=isolated_module('logger',{'gateway_logger','logger'})
    logger.get_log_dir=lambda:root/'logs'
    logger.gateway_logger=logger.GatewayLogger();logger.logger=logger.gateway_logger.get_logger()
    import main
    from fastapi.testclient import TestClient
    import openpyxl
    cases=[]
    with TestClient(main.app) as client:
        def capture(method,path,body=None):
            kwargs={} if body is None else {'json':body}
            response=client.request(method,path,**kwargs)
            cases.append({'method':method,'path':path,'body':body,'status':response.status_code,'expected':response.json()})
        capture('GET','/api/health')
        capture('GET','/api/config')
        capture('GET','/api/devices')
        capture('GET','/api/mqtt')
        capture('GET','/api/ai/config')
        capture('GET','/api/collect/status')
        capture('POST','/api/devices',{'id':'d1','name':'温度计','protocol':'modbus_tcp','points':[{'name':'温度','address':'40001'}]})
        capture('POST','/api/devices',{'id':'rtu','name':'串口','protocol':'modbus_rtu','port':'COM9','parity':None})
        capture('GET','/api/devices')
        capture('PUT','/api/devices/d1',{'id':'ignored','name':'改名','protocol':'s7'})
        capture('PUT','/api/devices/missing',{'id':'x','name':'missing','protocol':'s7'})
        capture('DELETE','/api/devices/missing')
        capture('GET','/api/devices')
        capture('DELETE','/api/devices/rtu')
        capture('POST','/api/devices/test',{'device':{'id':'x','name':'x','protocol':'unknown'}})
        capture('PUT','/api/mqtt',{'host':'localhost','enabled':False})
        capture('GET','/api/mqtt')
        capture('PUT','/api/ai/config',{'model':'mock-model'})
        capture('GET','/api/ai/config')
        capture('POST','/api/ai/parse',{'input':'西门子PLC'})
        capture('PUT','/api/config',{'gateway_name':'新网关','collect_interval':50,'unknown':'ignored','ai':{'enabled':True}})
        capture('GET','/api/config')
        capture('POST','/api/config/reload')
        capture('GET','/api/config')
        capture('POST','/api/logs/level?level=DEBUG')
        capture('POST','/api/logs/level?level=weird')
        # Original schemas and routes, for route coverage auditing.
        (output/'openapi.json').write_text(json.dumps(main.app.openapi(),ensure_ascii=False,indent=2),encoding='utf-8')
        excel=[]
        sheets={
            'points_cn':[['点位名称','地址','数据类型','倍率','偏移量','单位'],['温度','40001','float',.1,-2,'℃'],['开关','00001','BOOL',1,0,''],['错误类型','D10','abc','bad',None,'rpm'],[None,None,None,None,None,None],['缺地址',None]],
            'points_en':[['NAME','ADDRESS','TYPE','RATE','OFFSET','UNIT'],['A',40001,'uint16',2,1,'bar'],['B','DB1.DBD0',None,None,None,None]],
            'missing_columns':[['姓名','寄存器'],['A','40001']],
            'empty_rows':[['name','address'],['',None]],
            'header_only':[['name','address']],
            'leading_empty_row':[[None,None],['name','address'],['A','40001']],
        }
        for name,rows in sheets.items():
            workbook=openpyxl.Workbook();sheet=workbook.active
            for row in rows:sheet.append(row)
            path=output/f'{name}.xlsx';workbook.save(path)
            response=client.post('/api/devices/import-excel',files={'file':(path.name,path.read_bytes())})
            excel.append({'file':path.name,'expected':response.json()})
        workbook=openpyxl.Workbook();workbook.active.append(['wrong']);sheet=workbook.create_sheet('Active');sheet.append(['name','address']);sheet.append(['active-point','40001']);workbook.active=1
        path=output/'active_sheet.xlsx';workbook.save(path)
        excel.append({'file':path.name,'expected':client.post('/api/devices/import-excel',files={'file':(path.name,path.read_bytes())}).json()})
        (output/'excel_oracle.json').write_text(json.dumps(excel,ensure_ascii=False,indent=2),encoding='utf-8')
    (output/'api_oracle.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2),encoding='utf-8')
    for handler in logger.logger.handlers[:]: handler.close();logger.logger.removeHandler(handler)
print(f'Captured {len(cases)} API transactions and {len(excel)} workbooks')
